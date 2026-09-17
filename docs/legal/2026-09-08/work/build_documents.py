from pathlib import Path
import re, json
from datetime import datetime, timezone
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
OUT.mkdir(exist_ok=True)
FILENAMES = {
 '00_readiness_and_rkn': '00_Инструкция_и_сведения_для_РКН.docx',
 '01_user_agreement': '01_Пользовательское_соглашение.docx',
 '02_privacy_policy': '02_Политика_обработки_данных.docx',
 '03_optional_processing_consents': '03_Согласия_на_обработку_данных.docx',
 '04_dissemination_consent': '04_Согласие_на_распространение.docx',
 '05_recommendation_rules': '05_Правила_рекомендаций.docx',
 '06_internal_data_procedure': '06_Внутренний_порядок_обработки_данных.docx',
 '07_request_and_response_forms': '07_Формы_обращений_и_ответов.docx',
}

def font(style, size, bold=False):
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    rp = style.element.get_or_add_rPr()
    for ch in list(rp):
        if ch.tag == qn('w:color'):
            for key in list(ch.attrib):
                if 'theme' in key:
                    del ch.attrib[key]
    fs = rp.find(qn('w:rFonts'))
    if fs is not None:
        for key in list(fs.attrib):
            if 'Theme' in key or 'theme' in key:
                del fs.attrib[key]
        fs.set(qn('w:ascii'),'Arial')
        fs.set(qn('w:hAnsi'),'Arial')
        fs.set(qn('w:cs'),'Arial')

def add_link(p, text, url):
    rel = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    r = OxmlElement('w:r'); rp = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'),'1F4D70'); rp.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'),'single'); rp.append(u)
    r.append(rp); t = OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); p._p.append(h)

def add_text(p, text):
    for chunk in re.split(r'(\[\[.*?\]\]|https?://[^\s]+)',text):
        if not chunk: continue
        if chunk.startswith('[['):
            r=p.add_run(chunk.replace('_',' '))
            r.font.highlight_color=WD_COLOR_INDEX.YELLOW
        elif chunk.startswith(('https://','http://')):
            tail=''
            while chunk.endswith(('.',',',';')):
                tail=chunk[-1]+tail; chunk=chunk[:-1]
            add_link(p,chunk,chunk)
            if tail:p.add_run(tail)
        else:p.add_run(chunk)

def page_field(p,code):
    r=p.add_run(); a=OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'begin');r._r.append(a)
    r=p.add_run();a=OxmlElement('w:instrText');a.set(qn('xml:space'),'preserve');a.text=' '+code+' ';r._r.append(a)
    r=p.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'end');r._r.append(a)

def new_doc(title):
    d=Document(); s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(1.9);s.bottom_margin=Cm(1.9);s.left_margin=Cm(2.1);s.right_margin=Cm(2.1)
    s.header_distance=Cm(.8);s.footer_distance=Cm(.8)
    for name,size,bold in [('Normal',10.5,False),('Title',21,True),('Subtitle',10.5,False),('Heading 1',13,True),('Heading 2',11,True),('Heading 3',10.5,True),('List Bullet',10.5,False),('Header',8,False),('Footer',8,False)]:
        font(d.styles[name],size,bold)
    n=d.styles['Normal'].paragraph_format;n.line_spacing=1.08;n.space_after=Pt(6);n.widow_control=True
    for name in ['Heading 1','Heading 2','Heading 3']:
        pf=d.styles[name].paragraph_format;pf.space_before=Pt(12);pf.space_after=Pt(5);pf.keep_with_next=True
    d.styles['Title'].paragraph_format.space_after=Pt(10)
    d.styles['Title'].paragraph_format.keep_with_next=True
    for style in d.styles:
        pp=style.element.find(qn('w:pPr'))
        if pp is not None:
            for el in list(pp):
                if el.tag==qn('w:pBdr'):pp.remove(el)
    header=s.header.paragraphs[0];header.text='НаТреню  |  Комплект документов  |  Проект 08.09.2026'
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Проект для заполнения  •  ');page_field(footer,'PAGE');footer.add_run(' / ');page_field(footer,'NUMPAGES')
    d.core_properties.title=title
    d.core_properties.subject='Бесплатное приложение НаТреню без рекламы'
    d.core_properties.author='НаТреню'
    d.core_properties.keywords='Проект; персональные данные; самозанятый; НаТреню'
    d.core_properties.created=datetime(2026,9,8,0,0,tzinfo=timezone.utc)
    d.core_properties.modified=datetime(2026,9,8,0,0,tzinfo=timezone.utc)
    settings=d.settings.element
    lang=OxmlElement('w:themeFontLang');lang.set(qn('w:val'),'ru-RU');settings.append(lang)
    return d

def add_table(d,lines):
    rows=[[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
    rows=[r for r in rows if not all(re.fullmatch(r'[-: ]+',c) for c in r)]
    cols=len(rows[0]);t=d.add_table(rows=0,cols=cols);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    widths=[5.0,5.9,5.9] if cols==3 else [5.2,11.6]
    if cols not in [2,3]:widths=[16.8/cols]*cols
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        pr=t.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:
            rep=OxmlElement('w:tblHeader');pr.append(rep)
        for ci,(cell,content) in enumerate(zip(cells,row)):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side,val in [('top','75'),('bottom','75'),('left','90'),('right','90')]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');m.append(el)
            cp.append(m)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'404040' if ri==0 else ('F4F4F4' if ri%2==0 else 'FFFFFF'));cp.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.03
            add_text(p,content)
            for run in p.runs:
                run.font.size=Pt(9)
                if ri==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=0.3

def build(src):
    lines=src.read_text().splitlines();title=lines[0][2:];d=new_doc(title)
    if src.stem in ['01_user_agreement', '05_recommendation_rules']:
        d.styles['Normal'].paragraph_format.line_spacing = 1.04
        d.styles['Normal'].paragraph_format.space_after = Pt(5)
        for name in ['Heading 1', 'Heading 2', 'Heading 3']:
            d.styles[name].paragraph_format.space_before = Pt(10)
    if src.stem == '01_user_agreement':
        d.styles['Normal'].paragraph_format.line_spacing = 1.0
        d.styles['Normal'].paragraph_format.space_after = Pt(4)
    if src.stem == '02_privacy_policy':
        d.styles['Normal'].paragraph_format.line_spacing = 1.0
        d.styles['Normal'].paragraph_format.space_after = Pt(4)
        for name in ['Heading 1', 'Heading 2', 'Heading 3']:
            d.styles[name].paragraph_format.space_before = Pt(10)
    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line=='[[PAGEBREAK]]':
            p=d.add_paragraph();p.paragraph_format.space_after=Pt(0);p.add_run().add_break(WD_BREAK.PAGE);i+=1;continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
            add_table(d,block);continue
        if line.startswith('# '):p=d.add_paragraph(line[2:],style='Title')
        elif line.startswith('### '):p=d.add_paragraph(line[4:],style='Heading 2')
        elif line.startswith('## '):p=d.add_paragraph(line[3:],style='Heading 1')
        elif line.startswith('- '):
            content=line[2:];p=d.add_paragraph(style='Normal' if content.startswith('□') else 'List Bullet')
            match=re.fullmatch(r'(.+?): (https?://\S+)',content)
            if match:add_link(p,match.group(1),match.group(2))
            else:add_text(p,content)
        else:p=d.add_paragraph();add_text(p,line)
        i+=1
    path=OUT/FILENAMES[src.stem];d.save(path)
    return {'file':str(path),'title':title,'source':str(src),'paragraphs':len(d.paragraphs),'tables':len(d.tables)}

if __name__=='__main__':
    manifest=[build(src) for src in sorted((ROOT/'source').glob('*.md'))]
    (ROOT/'work'/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    print(json.dumps(manifest,ensure_ascii=False,indent=2))
